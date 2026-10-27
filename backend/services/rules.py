from __future__ import annotations

import json
import re
from io import BytesIO

from openpyxl import Workbook, load_workbook
from sqlalchemy.orm import Session

from errors import AppError, conflict, not_found, unprocessable
from schemas.time_rules import validate_published_time_params
from models import (
    CheckRule,
    ChecklistItem,
    MaterialType,
    PoliticalTerm,
    PromptTemplate,
    RuleVersion,
    SensitivePhrase,
    SuggestionTemplate,
    TimeNodeDef,
    TimeRule,
)
from seed_rules import snapshot_rules, lock_rule_publication
from services.audit import log
from utils import now

KIND_MAP = {
    "checklists": (ChecklistItem, "清单"),
    "material-types": (MaterialType, "材料类型"),
    "time-nodes": (TimeNodeDef, "时间节点"),
    "time-rules": (TimeRule, "时间规则"),
    "check-rules": (CheckRule, "校验规则"),
    "terms": (PoliticalTerm, "术语"),
    "sensitive-phrases": (SensitivePhrase, "敏感表述"),
    "suggestion-templates": (SuggestionTemplate, "建议模板"),
    "prompt-templates": (PromptTemplate, "提示词"),
}


def list_checklists(db: Session, stage_code=None, q=None, is_active=None):
    items = db.query(ChecklistItem).order_by(ChecklistItem.stage_code, ChecklistItem.sort_order).all()
    if stage_code:
        items = [i for i in items if i.stage_code == int(stage_code)]
    if is_active is not None:
        flag = str(is_active).lower() in ("1", "true", "yes")
        items = [i for i in items if i.is_active == flag]
    grouped = {}
    from utils import STAGE_NAMES, stage_name
    for it in items:
        grouped.setdefault(it.stage_code, []).append(_checklist_dict(it))
    return [{"stage_code": k, "stage_name": stage_name(db, k), "items": v} for k, v in grouped.items()]


def _checklist_dict(it: ChecklistItem) -> dict:
    mt = it.material_type
    return {
        "id": it.id, "stage_code": it.stage_code, "material_type_id": it.material_type_id,
        "material_type": {"id": mt.id, "code": mt.code, "name": mt.name} if mt else None,
        "quantity_mode": it.quantity_mode, "fixed_count": it.fixed_count, "period_months": it.period_months,
        "period_start_node": it.period_start_node, "period_end_node": it.period_end_node,
        "is_required": it.is_required, "requirement_text": it.requirement_text,
        "sort_order": it.sort_order, "is_active": it.is_active,
    }


def orm_dict(obj, extra=None) -> dict:
    data = {}
    for c in obj.__table__.columns:
        v = getattr(obj, c.name)
        if hasattr(v, "isoformat"):
            v = v.isoformat()
        data[c.name] = v
    if extra:
        data.update(extra)
    if isinstance(obj, SuggestionTemplate):
        data["variables"] = re.findall(r"\{([^}]+)\}", obj.template_text or "")
    return data


def _audit_rule(db, actor, action, kind, object_id, summary, version_no):
    log(db, actor, f"RULE_{action}", kind, object_id, {"summary": summary, "rule_version_no": version_no})


def create_and_snapshot(db, actor, kind, obj, summary) -> tuple:
    lock_rule_publication(db)
    if kind == "time-rules":
        obj.params = validate_published_time_params(db, obj.params)
    db.add(obj)
    db.flush()
    ver = snapshot_rules(db, kind, "CREATE", summary, obj.id if hasattr(obj, "id") else None, actor.id)
    _audit_rule(db, actor, "CREATE", kind, getattr(obj, "id", None), summary, ver.version_no)
    return obj, ver.version_no


def update_and_snapshot(db, actor, kind, obj, fields, summary):
    _lock_and_refresh(db, obj)
    if kind == "time-rules" and "params" in fields:
        fields = {**fields, "params": validate_published_time_params(db, fields["params"])}
    elif kind == "time-rules" and fields.get("is_active") is True:
        validate_published_time_params(db, obj.params)
    for k, v in fields.items():
        may_clear = kind == "time-nodes" and k in ("stage_code", "required_from_stage", "derive_expr")
        if (v is not None or may_clear) and hasattr(obj, k) and k not in ("id", "code", "created_at"):
            setattr(obj, k, v)
    if hasattr(obj, "updated_at"):
        obj.updated_at = now()
    db.flush()
    ver = snapshot_rules(db, kind, "UPDATE", summary, getattr(obj, "id", None), actor.id)
    _audit_rule(db, actor, "UPDATE", kind, getattr(obj, "id", None), summary, ver.version_no)
    return obj, ver.version_no


def delete_and_snapshot(db, actor, kind, obj, summary):
    _lock_and_refresh(db, obj)
    oid = getattr(obj, "id", None)
    try:
        db.delete(obj)
        db.flush()
    except Exception:
        raise conflict("条目已被引用，请改为停用")
    ver = snapshot_rules(db, kind, "DELETE", summary, oid, actor.id)
    _audit_rule(db, actor, "DELETE", kind, oid, summary, ver.version_no)
    return ver.version_no


def _lock_and_refresh(db, obj):
    from sqlalchemy.exc import InvalidRequestError
    lock_rule_publication(db)
    try:
        db.refresh(obj)
    except InvalidRequestError:
        raise not_found('规则条目已被删除，请刷新后重试')


def export_kind(db: Session, kind: str) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = kind[:20]
    model = KIND_MAP.get(kind, (None, None))[0]
    if model is None:
        raise unprocessable("不支持导出该类型")
    rows = db.query(model).all()
    if not rows:
        ws.append(["empty"])
        return wb
    cols = [c.name for c in model.__table__.columns]
    ws.append(cols)
    for r in rows:
        ws.append([_cell(getattr(r, c)) for c in cols])
    return wb


def _cell(v):
    if v is None:
        return None
    if hasattr(v, "isoformat"):
        return v.isoformat()
    if isinstance(v, (dict, list)):
        return json.dumps(v, ensure_ascii=False)
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@"):
        return "'" + v
    return v


def import_kind(db, actor, kind, data: bytes, mode="upsert", dry_run=True) -> dict:
    model = KIND_MAP.get(kind, (None, None))[0]
    if model is None:
        raise unprocessable("不支持导入该类型")
    if mode not in ("append", "upsert", "replace"):
        raise unprocessable("不支持的导入模式")
    wb = load_workbook(BytesIO(data), data_only=True)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    created = updated = skipped = 0
    errors = []
    total = 0
    unique_field = "code" if hasattr(model, "code") else ("standard_text" if hasattr(model, "standard_text") else "phrase")
    if not dry_run:
        lock_rule_publication(db)
    records = []
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not any(row):
            continue
        total += 1
        rec = dict(zip(headers, row))
        rec = {k: _parse_cell(v) for k, v in rec.items() if k}
        key = rec.get(unique_field)
        if not key:
            errors.append({"row": i, "message": f"缺少 {unique_field}"})
            continue
        existing = db.query(model).filter(getattr(model, unique_field) == key).first() if mode != "replace" and hasattr(model, unique_field) else None
        if existing and mode == "append":
            skipped += 1
            continue
        if kind == "time-rules":
            try:
                rec["params"] = validate_published_time_params(db, rec.get("params", existing.params if existing else None))
            except AppError as exc:
                errors.append({"row": i, "message": exc.biz_message, "details": exc.biz_data})
                continue
        records.append((i, rec, existing))
        if dry_run:
            if existing:
                updated += 1
            else:
                created += 1
            continue
    if errors and not dry_run:
        raise unprocessable("导入存在错误，未保存任何修改", errors)
    if mode == "replace" and not dry_run:
        db.query(model).delete()
        db.flush()
    for i, rec, existing in records if not dry_run else []:
        try:
            if existing and mode in ("upsert", "replace"):
                for k, v in rec.items():
                    if hasattr(existing, k) and k not in ("id",):
                        setattr(existing, k, v)
                updated += 1
            elif existing and mode == "append":
                skipped += 1
            else:
                rec.pop("id", None)
                obj = model(**{k: v for k, v in rec.items() if hasattr(model, k)})
                if hasattr(obj, "created_at") and not obj.created_at:
                    obj.created_at = now()
                if hasattr(obj, "updated_at"):
                    obj.updated_at = now()
                db.add(obj)
                created += 1
        except Exception as e:
            errors.append({"row": i, "message": str(e)})
    if errors and not dry_run:
        # 任一行出错则整体回滚（含 replace 模式的清表），保证全有或全无
        db.rollback()
        raise unprocessable("导入存在错误，已整体回滚", errors)
    ver_no = None
    if not dry_run:
        ver = snapshot_rules(db, kind, "IMPORT", f"导入 {kind} {mode}", None, actor.id)
        ver_no = ver.version_no
        _audit_rule(db, actor, "IMPORT", kind, None, f"导入 {kind} {mode}: 新增 {created} 更新 {updated}", ver_no)
    return {"total": total, "created": created, "updated": updated, "skipped": skipped, "errors": errors, "rule_version_no": ver_no}


def _parse_cell(v):
    if isinstance(v, str) and (v.startswith("{") or v.startswith("[")):
        try:
            return json.loads(v)
        except json.JSONDecodeError:
            return v
    return v


def version_diff(a: dict, b: dict) -> list:
    changes = []
    for kind, rows in (b or {}).items():
        old_rows = {str(x.get("code") or x.get("id")): x for x in (a or {}).get(kind) or []}
        new_rows = {str(x.get("code") or x.get("id")): x for x in rows or []}
        for k, nv in new_rows.items():
            ov = old_rows.get(k)
            if not ov:
                changes.append({"kind": kind, "code": k, "field": "*", "old": None, "new": "created"})
                continue
            for f, val in nv.items():
                if ov.get(f) != val:
                    changes.append({"kind": kind, "code": k, "field": f, "old": ov.get(f), "new": val})
        for k in old_rows:
            if k not in new_rows:
                changes.append({"kind": kind, "code": k, "field": "*", "old": "exists", "new": None})
    return changes
