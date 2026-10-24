import fitz
import pytest
from docx import Document

from ai import parser


@pytest.fixture
def storage(tmp_path, monkeypatch):
    monkeypatch.setattr(parser, "STORAGE_DIR", tmp_path)
    monkeypatch.setattr(parser, "ensure_storage", lambda: tmp_path)
    (tmp_path / "materials").mkdir()
    return tmp_path


def test_long_docx_preserves_all_paragraphs_and_real_page_coordinates(storage):
    doc = Document()
    for i in range(125):
        doc.add_paragraph(f"段落 P{i:03d}：请根据实际学习和实践情况完善材料。")
    source = storage / "materials/long.docx"
    doc.save(source)

    result = parser.parse_version_file("materials/long.docx", "docx")

    with fitz.open(storage / result["preview_pdf_path"]) as preview:
        assert len(preview) > 1
        assert result["page_count"] == len(preview) == len(result["pages"])
        text = "".join("".join(page.get_text() for page in preview).split())
        for i in range(125):
            assert f"P{i:03d}" in text
        assert result["pages"][-1]["page_no"] == len(preview)
        assert all(block["bbox"] and len(block["bbox"]) == 4
                   for page in result["pages"] for block in page["blocks"])
        for page in preview:
            spans = [s for b in page.get_text("dict")["blocks"] if "lines" in b
                     for line in b["lines"] for s in line["spans"]]
            assert min(s["size"] for s in spans) >= 10.5
    png, page_count, _, _ = parser.render_page_png(result["preview_pdf_path"], result["page_count"])
    assert png.startswith(b"\x89PNG") and page_count == result["page_count"]


def test_paragraphs_and_tables_keep_their_original_order(storage):
    doc = Document()
    doc.add_paragraph("BEFORE_TABLE")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "TABLE_VALUE"
    doc.add_paragraph("AFTER_TABLE <原文> & 说明")
    doc.save(storage / "materials/table.docx")

    result = parser.parse_version_file("materials/table.docx", "docx")

    for text in (result["full_text"], "\n".join(p["text"] for p in result["pages"])):
        assert text.index("BEFORE_TABLE") < text.index("TABLE_VALUE") < text.index("AFTER_TABLE")
        assert "<原文>" in text and "&" in text


def test_preview_failure_is_reported_instead_of_writing_a_fake_pdf(storage, monkeypatch):
    def fail(*args, **kwargs):
        raise RuntimeError("conversion failed")
    monkeypatch.setattr(fitz, "Story", fail)

    with pytest.raises(RuntimeError, match="conversion failed"):
        parser._docx_to_preview_pdf(storage / "materials/broken.docx", "正文")

    assert not (storage / "previews/broken.pdf").exists()
