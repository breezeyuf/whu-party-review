from __future__ import annotations

from html import escape
from io import BytesIO
from pathlib import Path
from urllib.parse import quote

from fastapi.responses import Response
from sqlalchemy.orm import Session

from models import ExportFile, Issue, Notice
from services import storage as store
from utils import now, get_setting
from datetime import timedelta


def run_export_job(db: Session, job) -> None:
    params = job.params or {}
    export_type = params.get("export_type") or "STATS_REPORT"
    filters = params.get("filters") or {}
    # 服务端注入的作用域优先于客户端 filters，防止越权导出
    scope_college = params.get("scope_college_id")
    target_college = scope_college or filters.get("college_id")
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = export_type[:20]
    row_count = 0
    if export_type == "AI_QUALITY":
        from models import Submission
        ws.append(["rule_code", "title", "confirmed", "rejected", "modified"])
        q = (
            db.query(Issue.rule_code, Issue.review_status)
            .join(Submission, Submission.id == Issue.submission_id)
            .filter(Issue.source != "MANUAL")
        )
        if target_college:
            q = q.filter(Submission.college_id == int(target_college))
        if filters.get("batch_id"):
            q = q.filter(Submission.batch_id == int(filters["batch_id"]))
        rows = q.all()
        from collections import Counter, defaultdict
        c = defaultdict(Counter)
        for code, st in rows:
            c[code or ""][st] += 1
        for code, ct in c.items():
            ws.append([_safe_cell(code), "", ct.get("CONFIRMED", 0), ct.get("REJECTED", 0), ct.get("MODIFIED", 0)])
            row_count += 1
    elif export_type == "DESENSITIZED_DATA":
        from models import Archive
        ws.append(["anon_code", "stage", "status", "is_demo"])
        q = db.query(Archive)
        if filters.get("batch_id"):
            q = q.filter(Archive.batch_id == int(filters["batch_id"]))
        if target_college:
            q = q.filter(Archive.college_id == int(target_college))
        for a in q.all():
            ws.append([a.anon_code, a.current_stage, a.status, a.is_demo])
            row_count += 1
    else:
        from models import Submission
        ws.append(["submission_id", "stage", "status", "round_no"])
        q = db.query(Submission)
        if target_college:
            q = q.filter(Submission.college_id == int(target_college))
        if filters.get("batch_id"):
            q = q.filter(Submission.batch_id == int(filters["batch_id"]))
        for s in q.limit(5000).all():
            ws.append([s.id, s.stage_code, s.status, s.round_no])
            row_count += 1
    buf = BytesIO()
    wb.save(buf)
    data = buf.getvalue()
    rel = f"exports/{job.id}_{export_type}.xlsx"
    path = store.abs_path(rel)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(data)
    days = int(get_setting(db, "retention", {}).get("export_days") or 7)
    ef = ExportFile(
        export_type=export_type, params=filters, file_name=f"{export_type}.xlsx",
        stored_path=rel, row_count=row_count, created_by=job.created_by,
        created_at=now(), expires_at=now() + timedelta(days=days),
    )
    db.add(ef)
    db.flush()
    job.result = {"export_id": ef.id, "file_name": ef.file_name, "row_count": row_count,
                  "download_url": f"/api/exports/{ef.id}/download"}


def _safe_cell(v):
    """规避 xlsx 公式注入：以 = + - @ 开头的文本加前缀。"""
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@"):
        return "'" + v
    return v


def xlsx_response(filename: str, wb) -> Response:
    buf = BytesIO()
    wb.save(buf)
    data = buf.getvalue()
    return Response(
        content=data,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )


def export_notice_file(notice: Notice, fmt: str = "pdf") -> Response:
    student = (notice.content or {}).get("student") or {}
    anon = student.get("student_no") or str(notice.id)
    title = notice.title
    items = (notice.content or {}).get("items") or []
    comps = (notice.content or {}).get("completion_items") or []
    note = (notice.content or {}).get("note") or ""
    if fmt == "docx":
        from docx import Document
        doc = Document()
        doc.add_heading(title, 0)
        intro = (notice.content or {}).get("intro") or ""
        doc.add_paragraph(intro)
        for it in items:
            doc.add_paragraph(f"{it.get('no')}. [{it.get('severity')}] {it.get('problem')}")
            if it.get("suggestion"):
                doc.add_paragraph(f"修改建议：{it.get('suggestion')}")
        for it in comps:
            doc.add_paragraph(f"缺项：{it.get('material')} 缺 {it.get('missing')} 份")
        if note:
            doc.add_paragraph(f"审核意见：{note}")
        buf = BytesIO()
        doc.save(buf)
        name = f"整改通知_{anon}_第{notice.round_no}轮.docx"
        return Response(
            content=buf.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(name)}"},
        )
    import fitz

    def html_text(value):
        # 问题原文中的标签和换行属于正文，不能作为排版 HTML 执行。
        text = "" if value is None else str(value)
        return escape(text).replace("\r\n", "\n").replace("\r", "\n").replace("\n", "<br/>")

    parts = [f"<h2>{html_text(title)}</h2>", f"<p>{html_text((notice.content or {}).get('intro'))}</p>"]
    for it in items:
        parts.append(
            f"<p>{html_text(it.get('no'))}. [{html_text(it.get('severity'))}] "
            f"{html_text(it.get('problem'))}<br/>建议：{html_text(it.get('suggestion'))}</p>"
        )
    for it in comps:
        parts.append(f"<p>缺项：{html_text(it.get('material'))} 缺 {html_text(it.get('missing'))} 份</p>")
    if note:
        parts.append(f"<p>审核意见：{html_text(note)}</p>")

    # Story 将未排完的正文续写到下一页，保持字号，避免单页缩放到无法阅读。
    story = fitz.Story(
        html="".join(parts), em=11,
        user_css="body { font-size: 11pt; line-height: 1.5; margin: 0; } "
                 "h2 { font-size: 16pt; margin: 0 0 14pt; } p { margin: 0 0 10pt; }",
    )
    page_rect = fitz.paper_rect("a4")
    content_rect = page_rect + (42, 42, -42, -48)
    buf = BytesIO()
    with fitz.DocumentWriter(buf) as writer:
        story.write(writer, lambda _index, _filled: (page_rect, content_rect, None))
    with fitz.open(stream=buf.getvalue(), filetype="pdf") as doc:
        for page in doc:
            label = f"{page.number + 1} / {len(doc)}"
            label_width = fitz.get_text_length(label, fontname="helv", fontsize=9)
            page.insert_text(((page.rect.width - label_width) / 2, page.rect.height - 24), label, fontsize=9)
        data = doc.tobytes(garbage=4, deflate=True)
    name = f"整改通知_{anon}_第{notice.round_no}轮.pdf"
    return Response(
        content=data, media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(name)}"},
    )
