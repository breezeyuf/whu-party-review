from types import SimpleNamespace as N
import pytest
from ai.checks.format import run_format
from ai.checks.content import run_content


def fixture(code, params=None, text='正文', features=None):
    r = N(code=code, category='FORMAT' if code.startswith('R-F') else 'CONTENT', engine='RULE', params=params or {},
          applies_to=[], is_active=True, name=code, severity='WARN', needs_human_judgement=False, basis='', message_template='')
    m = N(id=1, status='ACTIVE', type_id=1, title='材料')
    p = N(corrected_text=None, full_text=text, corrected_fields=None, extracted_fields={}, char_count=len(text),
          pages=[], visual_marks={}, format_features=features or {})
    t = N(code='AUTOBIOGRAPHY', format_requirements={}, element_template=[])
    rs = N(rows=lambda group: [r] if group == 'check_rules' else [])
    return r, m, p, t, rs


def check(code, params=None, features=None, text='正文'):
    r,m,p,t,rs = fixture(code, params, text, features)
    args = (None, [m], {1:N(id=1)}, {1:p}, {1:t})
    return run_format(*args, rule_set=rs) if r.category == 'FORMAT' else run_content(None,N(archive=None),*args[1:],False,rule_set=rs)


@pytest.mark.parametrize('code', ['R-F09','R-C09','R-F99','R-C99'])
def test_missing_rule_implementation_or_parameters_never_pass(code):
    assert check(code)[1][0]['result'] == 'SKIPPED'


@pytest.mark.parametrize('field,actual,expected', [('size_pt',9,12),('line_spacing',1,1.5)])
def test_layout_checks_more_than_font(field,actual,expected):
    assert check('R-F08',{'docx_layout':{field:expected}},{field:actual})[1][0]['result']=='FAIL'


def test_unmeasured_layout_is_skipped():
    assert check('R-F08',{'docx_layout':{'size_pt':12}},{'font':'宋体'})[1][0]['result']=='SKIPPED'


def test_orientation_and_binding_are_checked():
    features={'sections':[{'orientation':'landscape','margins_cm':{'left':1}}]}
    assert check('R-F09',{'orientation':'portrait','min_binding_margin_cm':2.5},features)[1][0]['result']=='FAIL'
    assert check('R-F09',{'orientation':'landscape','min_binding_margin_cm':0.5},features)[1][0]['result']=='PASS'


def test_quotes_require_configured_authoritative_reference():
    params={'references':[{'title':'测试规范','text':'正确的原文','source':'测试来源'}]}
    assert check('R-C09',params,text='《测试规范》指出：“错误的引文”')[1][0]['result']=='FAIL'
    assert check('R-C09',params,text='《测试规范》指出：“正确的原文”')[1][0]['result']=='PASS'
    assert check('R-C09',params,text='无引用')[1][0]['result']=='SKIPPED'


@pytest.mark.parametrize('params', [{'references':'bad'}, {'references':[None,42,{'title':3}]}])
def test_invalid_quote_configuration_is_skipped(params):
    assert check('R-C09',params)[1][0]['result']=='SKIPPED'


@pytest.mark.parametrize('threshold', ['bad', float('nan'), -1, 2])
def test_invalid_similarity_threshold_is_skipped(threshold):
    assert check('R-C10',{'threshold':threshold})[1][0]['result']=='SKIPPED'


def test_docx_measures_inherited_styles_and_all_sections():
    from docx import Document
    from docx.shared import Pt
    from ai.layout import docx_features, check_layout
    doc=Document(); style=doc.styles['Normal']; style.font.name='宋体'; style.font.size=Pt(12)
    style.paragraph_format.line_spacing=1.5
    doc.add_paragraph('正文'); doc.add_paragraph('另段').runs[0].font.size=Pt(9)
    f=docx_features(doc)
    assert f['samples'][0]['size_pt']==12 and f['samples'][1]['size_pt']==9
    assert len(f['sections'])==1
    problems, complete=check_layout(f,{'size_pt':12,'line_spacing':1.5})
    assert problems and complete


def test_corrected_text_length_and_deliberate_empty_override_ocr():
    r,m,p,t,rs=fixture('R-F06',{'min_chars':{'AUTOBIOGRAPHY':3000}},text='旧'*10)
    p.corrected_text='新'*4000
    args=(None,[m],{1:N(id=1)},{1:p},{1:t})
    assert run_format(*args,rule_set=rs)[1][0]['result']=='PASS'
    p.corrected_text=''
    p.full_text='旧'*4000
    assert run_format(*args,rule_set=rs)[1][0]['result']=='FAIL'
    assert p.char_count==10  # Preserve original evidence, compute effective count.


def test_quality_uses_configured_dpi_and_threshold(monkeypatch):
    import fitz
    from ai.ocr import settings
    from ai.vision import page_blur
    with fitz.open() as doc:
        page=doc.new_page(width=100,height=100)
        token=settings.set({'min_dpi':200,'blur_threshold':1000})
        try:
            quality=page_blur(page)
            assert quality['dpi']==200 and quality['blur_threshold']==1000 and quality['is_blurry']
        finally: settings.reset(token)


def test_scanned_pdf_quality_and_ocr_failure_flag(monkeypatch,tmp_path):
    import fitz
    from ai import parser
    monkeypatch.setattr(parser,'STORAGE_DIR',tmp_path)
    monkeypatch.setattr(parser,'_ocr_page',lambda page:('',[],None,'OCR 超时'))
    source=tmp_path/'scan.pdf'
    with fitz.open() as doc:
        doc.new_page(width=100,height=100); doc.save(source)
    result=parser._parse_pdf(source)
    assert result['ocr_failed'] and result['pages'][0]['quality']['blur_score'] is not None


def test_visual_uses_ocr_positions_for_multiple_signatures():
    import fitz
    from ai.vision import detect_visual
    with fitz.open() as doc:
        page=doc.new_page(width=300,height=300)
        page.draw_line((60,40),(95,50),width=3)
        page.draw_line((60,140),(95,150),width=3)
        pages=[{'page_no':1,'blocks':[{'text':'本人签字','bbox':[10,30,50,50]}, {'text':'介绍人','bbox':[10,130,50,150]}]}]
        found=detect_visual(doc,pages)['signatures']
        assert len(found)==2 and found[0]['bbox'][1]<100 and found[1]['bbox'][1]>100


def test_one_signature_does_not_satisfy_multiple_required_slots():
    from ai.checks.format import _check_one
    r,m,p,t,_=fixture('R-F03')
    t.format_requirements={'signature_slots':[{'slot':'本人','keywords':['签名']},{'slot':'介绍人','keywords':['签名']}]}
    p.extracted_fields={'signature_name':'某人'}
    p.visual_marks={'signatures':[{'slot':'签名','page_no':1,'bbox':[10,10,20,20],'present':True}]}
    issues=_check_one(r,m,N(id=1),p,t)
    assert len(issues)==2 and all(x.needs_human_judgement for x in issues)
    assert '未可靠定位' in issues[1].description


def test_missing_archive_fields_still_detect_cross_material_conflicts():
    from ai.checks.content import _apply_rule
    r,m,p,t,_=fixture('R-C06')
    p.extracted_fields={'ethnicity':'甲族','native_place':'甲地'}
    issues=_apply_rule(r,m,N(id=1),p,t,'',[],[],N(archive=N()),{}, {},
                      {1:p.extracted_fields,2:{'ethnicity':'乙族','native_place':'乙地'}})
    assert {i.location['field'] for i in issues}=={'ethnicity','native_place'}


def test_delivery_config_encrypts_and_never_sends(monkeypatch):
    from services import delivery
    from errors import AppError
    monkeypatch.setattr(delivery,'encrypt_field',lambda s:'enc:'+s)
    config=delivery.update_config({}, {'smtp_host':'mail.invalid','smtp_username':'test','smtp_sender':'test@invalid',
                                      'smtp_password':'private','sms_url':'https://sms.invalid/send','sms_api_key':'secret'})
    public=delivery.public_config(config)
    assert 'private' not in str(public) and 'secret' not in str(public)
    assert public['smtp_password_set'] and public['sms_api_key_set']
    assert delivery.check_config(config)['sent'] is False
    assert delivery.prepare_delivery('email','fake@invalid','test','test',config)['sent'] is False
    with pytest.raises(AppError): delivery.update_config(config,{'smtp_host':'elsewhere.invalid'})
    with pytest.raises(AppError): delivery.update_config(config,{'mode':'live'})
