from concurrent.futures import ThreadPoolExecutor
from datetime import date, timedelta
from io import BytesIO
from threading import Barrier
from types import SimpleNamespace
from uuid import uuid4

import msoffcrypto
import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy.orm import Session

from errors import AppError
from models import Archive, Batch, ImportPasswordFile, Org, TimeNodeExtraction, User, UserRole
from services.importer import import_users_xlsx
from services.import_passwords import consume_password_file, create_password_file
from utils import now
from tests.test_review_regressions import test_engine, case, as_client  # noqa: F401


def workbook(students=None, users=None):
    wb = Workbook()
    wb.remove(wb.active)
    for sheet, rows, headers in (("用户", users, ["username", "real_name", "gender", "org_code", "role", "phone", "email"]),
        ("学生档案", students, ["student_no", "real_name", "branch_code", "birth_date", "batch_name", "T02"])):
        if rows is not None:
            ws = wb.create_sheet(sheet)
            ws.append(headers)
            for row in rows:
                ws.append([row.get(key) for key in headers])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def prepare(c):
    c.branches[0].code, c.branches[1].code, c.college.code = "B0", "B1", "COL"
    c.db.flush()


def student(**values):
    return dict(student_no="new-2026", real_name="测试学生", branch_code="B0", **values)


@pytest.mark.parametrize("changes,column", [({"branch_code": "missing"}, "branch_code"),
    ({"branch_code": "COL"}, "branch_code"), ({"birth_date": "2026-02-30"}, "birth_date"),
    ({"T02": "2026-01-01garbage"}, "T02"), ({"T02": "=TODAY()"}, "T02"),
    ({"batch_name": "不存在"}, "batch_name")])
def test_dry_run_and_formal_reject_same_bad_student_rows_without_writes(case, changes, column):
    c = case
    prepare(c)
    row = student()
    row.update(changes)
    data = workbook(students=[row])
    before = c.db.query(User).count()
    dry = import_users_xlsx(c.db, c.org, data)
    assert dry["valid"] == 0
    assert any(e["column"] == column for e in dry["errors"])
    with pytest.raises(AppError) as exc:
        import_users_xlsx(c.db, c.org, data, dry_run=False)
    assert exc.value.status_code == 422
    assert c.db.query(User).count() == before
    assert c.db.query(ImportPasswordFile).count() == 0


def test_dry_run_checks_scope_and_inactive_organization(case):
    c = case
    prepare(c)
    other = c.add(Org, code="OTHER", org_type="COLLEGE", name="其他学院", parent_id=c.school.id,
                  created_at=now(), updated_at=now())
    c.branches[0].parent_id = other.id
    c.db.flush()
    result = import_users_xlsx(c.db, c.col, workbook(students=[student()]))
    assert result["valid"] == 0
    assert "无权" in result["errors"][0]["message"]
    c.branches[0].is_active = False
    c.db.flush()
    assert import_users_xlsx(c.db, c.org, workbook(students=[student()]))["valid"] == 0


def test_valid_dry_run_writes_nothing_and_formal_creates_archive_batch_and_nodes(case):
    c = case
    prepare(c)
    data = workbook(students=[student(birth_date=date(2003, 1, 1), batch_name=c.batches[0].name, T02="2026-01-01")])
    before = c.db.query(User).count()
    dry = import_users_xlsx(c.db, c.org, data)
    assert dry["valid"] == 1 and dry["errors"] == []
    assert c.db.query(User).count() == before
    assert "initial_passwords_file" not in dry
    result = import_users_xlsx(c.db, c.org, data, dry_run=False)
    assert (result["created_users"], result["created_archives"], result["imported_nodes"]) == (1, 1, 1)
    archive = c.db.query(Archive).filter_by(student_no="new-2026").one()
    assert archive.batch_id == c.batches[0].id
    assert archive.birth_date == date(2003, 1, 1)
    assert "new-2026" not in archive.anon_code
    assert c.db.query(TimeNodeExtraction).filter_by(archive_id=archive.id, node_code="T02").one().value_date == date(2026, 1, 1)
    assert "initial_passwords" not in result


def test_encrypted_workbook_requires_password_and_download_is_owner_only_once(case, as_client):
    c = case
    prepare(c)
    data = workbook(users=[dict(username="new-user", real_name="测试用户", org_code="B0", role="STUDENT")])
    result = import_users_xlsx(c.db, c.org, data, dry_run=False)
    info = result["initial_passwords_file"]
    record = c.db.get(ImportPasswordFile, info["id"])
    assert "new-user" not in record.encrypted_payload
    response = as_client(c.col).get(info["download_url"])
    assert response.status_code == 404
    assert record.consumed_at is None
    response = as_client(c.org).get(info["download_url"])
    assert response.status_code == 200
    assert response.headers["cache-control"] == "no-store, private"
    encrypted = msoffcrypto.OfficeFile(BytesIO(response.content))
    assert encrypted.is_encrypted()
    with pytest.raises(msoffcrypto.exceptions.InvalidKeyError):
        encrypted.load_key(password="wrong-password", verify_password=True)
    encrypted.load_key(password=info["password"], verify_password=True)
    plain = BytesIO()
    encrypted.decrypt(plain, verify_integrity=True)
    rows = list(load_workbook(plain, data_only=True).active.values)
    assert rows[1][0] == "new-user"
    from security import verify_password
    user = c.db.query(User).filter_by(username="new-user").one()
    assert verify_password(rows[1][1], user.password_hash)
    assert user.must_change_password
    assert record.encrypted_payload is None
    assert as_client(c.org).get(info["download_url"]).status_code == 410


def test_expired_file_is_denied_and_cleaned(case):
    c = case
    info = create_password_file(c.db, c.org, [("test", "synthetic-password")])
    record = c.db.get(ImportPasswordFile, info["id"])
    record.expires_at = now() - timedelta(seconds=1)
    c.db.flush()
    with pytest.raises(AppError) as exc:
        consume_password_file(c.db, c.org, record.id)
    assert exc.value.status_code == 410
    from worker import _scan
    _scan(c.db)
    assert c.db.query(ImportPasswordFile).count() == 0


def test_duplicate_rows_and_cross_sheet_identity_mismatch_rejected(case):
    c = case
    prepare(c)
    result = import_users_xlsx(c.db, c.org, workbook(students=[student(), student()]))
    assert any("重复" in e["message"] for e in result["errors"])
    result = import_users_xlsx(c.db, c.org, workbook(students=[student()], users=[
        dict(username="new-2026", real_name="其他姓名", org_code="B0", role="STUDENT")]))
    assert any("不一致" in e["message"] for e in result["errors"])


def test_user_and_archive_in_same_workbook_share_one_account(case):
    c = case
    prepare(c)
    data = workbook(students=[student()], users=[dict(username="new-2026", real_name="测试学生", org_code="B0", role="STUDENT")])
    result = import_users_xlsx(c.db, c.org, data, dry_run=False)
    assert result["created_users"] == result["created_archives"] == 1
    second = import_users_xlsx(c.db, c.org, data, dry_run=False)
    assert second["created_users"] == second["created_archives"] == 0
    assert "initial_passwords_file" not in second


def test_file_generation_failure_rolls_back_whole_import(case, monkeypatch):
    c = case
    prepare(c)
    from services import import_passwords
    def fail(*args):
        raise RuntimeError("synthetic encryption failure")
    monkeypatch.setattr(import_passwords, "create_password_file", fail)
    with pytest.raises(RuntimeError):
        import_users_xlsx(c.db, c.org, workbook(students=[student()]), dry_run=False)
    assert c.db.query(User).filter_by(username="new-2026").count() == 0
    assert c.db.query(Archive).filter_by(student_no="new-2026").count() == 0


def test_valid_row_before_invalid_row_still_writes_nothing(case):
    c = case
    prepare(c)
    good, bad = student(), student()
    bad.update(student_no="bad-2026", T02="not-a-date")
    with pytest.raises(AppError):
        import_users_xlsx(c.db, c.org, workbook(students=[good, bad]), dry_run=False)
    assert c.db.query(User).filter_by(username="new-2026").count() == 0


def test_college_cannot_grant_privileged_role_and_existing_teacher_cannot_be_linked(case):
    c = case
    prepare(c)
    result = import_users_xlsx(c.db, c.col, workbook(users=[
        dict(username="admin-new", real_name="测试", org_code="B0", role="SYS_ADMIN")]))
    assert result["valid"] == 0
    result = import_users_xlsx(c.db, c.org, workbook(students=[
        dict(student_no=c.reviewer.user.username, real_name=c.reviewer.user.real_name, branch_code="B0")]))
    assert any("不是该支部的学生" in e["message"] for e in result["errors"])


def test_import_api_returns_only_protected_file_metadata(case, as_client):
    c = case
    prepare(c)
    data = workbook(students=[student()])
    response = as_client(c.org).post("/api/users/import", data={"dry_run": "false"}, files={"file": ("a.xlsx", data)})
    assert response.status_code == 200, response.text
    assert response.headers["cache-control"] == "no-store, private"
    result = response.json()["data"]
    assert "initial_passwords" not in result
    assert result["initial_passwords_file"]["download_url"].startswith("/api/users/import-passwords/")


def test_batch_id_and_name_must_match_and_bad_upload_is_422(case, as_client):
    c = case
    prepare(c)
    data = workbook(students=[student(batch_name=c.batches[1].name)])
    result = import_users_xlsx(c.db, c.org, data, batch_id=c.batches[0].id)
    assert result["valid"] == 0
    client = as_client(c.org)
    for params, payload in (({"batch_id": "bad"}, data), ({}, b"not-excel")):
        response = client.post("/api/users/import", data=params, files={"file": ("a.xlsx", payload)})
        assert response.status_code == 422


def test_concurrent_downloads_only_one_succeeds(test_engine):
    with Session(test_engine) as db:
        user = User(username="concurrent-" + uuid4().hex, password_hash="unused", real_name="测试",
                    created_at=now(), updated_at=now())
        db.add(user)
        db.flush()
        actor = SimpleNamespace(id=user.id)
        info = create_password_file(db, actor, [(user.username, "synthetic-only")])
        db.commit()
        user_id = user.id
    barrier = Barrier(2)
    def download():
        with Session(test_engine) as db:
            barrier.wait(timeout=10)
            try:
                consume_password_file(db, SimpleNamespace(id=user_id), info["id"])
                db.commit()
                return 200
            except AppError as exc:
                db.rollback()
                return exc.status_code
    try:
        with ThreadPoolExecutor(max_workers=2) as pool:
            results = list(pool.map(lambda _: download(), range(2)))
        assert sorted(results) == [200, 410]
    finally:
        with Session(test_engine) as db:
            db.query(User).filter_by(id=user_id).delete()
            db.commit()
