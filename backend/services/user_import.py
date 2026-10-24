"""Plan and validate the entire workbook before creating any account."""
from datetime import date, datetime
from io import BytesIO
from uuid import uuid4
import re

from openpyxl import load_workbook
from sqlalchemy.exc import IntegrityError

from errors import unprocessable
from models import Archive, Batch, Org, Submission, TimeNodeDef, TimeNodeExtraction, User, UserRole
from security import encrypt_field, hash_password, mask_phone, random_password
from services.timeline import merge_archive_nodes
from services.users import PRIVILEGED, can_manage_user
from utils import now

ROLES = {"STUDENT", "BRANCH_REVIEWER", "CROSS_REVIEWER", "COLLEGE_REVIEWER", "ORG_ADMIN", "SYS_ADMIN"}


def _text(value):
    return str(value).strip() if value is not None else ""


def _date(value):
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if not isinstance(value, str) or not re.fullmatch(r"\d{4}-\d{2}-\d{2}", value.strip()):
        raise ValueError("请使用 Excel 日期或 YYYY-MM-DD")
    return date.fromisoformat(value.strip())


def import_users(db, actor, data, dry_run=True, batch_id=None):
    errors, user_plans, archive_plans = [], {}, []
    total = valid = 0

    def error(sheet, row, column, message):
        errors.append(dict(sheet=sheet, row=row, column=column, message=message))

    try:
        wb = load_workbook(BytesIO(data), data_only=False, read_only=True)
    except Exception:
        raise unprocessable("无法读取 Excel 文件，请使用有效的 .xlsx 模板")
    try:
        if not any(name in wb.sheetnames for name in ("用户", "学生档案")):
            raise unprocessable("缺少“用户”或“学生档案”工作表")
        selected_batch = None
        if batch_id not in (None, ""):
            try:
                if isinstance(batch_id, bool) or int(batch_id) <= 0:
                    raise ValueError()
                selected_batch = db.get(Batch, int(batch_id))
            except (TypeError, ValueError):
                selected_batch = None
            if not selected_batch:
                error("学生档案", 1, "batch_id", "指定批次不存在")
        node_defs = {n.code: n for n in db.query(TimeNodeDef).all()}
        seen = {"用户": set(), "学生档案": set()}
        for sheet in ("用户", "学生档案"):
            if sheet not in wb.sheetnames:
                continue
            rows = wb[sheet].iter_rows()
            header_row = next(rows, ())
            headers = [_text(c.value) for c in header_row]
            required = ("username", "real_name", "org_code") if sheet == "用户" else ("student_no", "real_name", "branch_code")
            if any(key not in headers for key in required) or len([h for h in headers if h]) != len(set(h for h in headers if h)):
                error(sheet, 1, "header", "表头缺失或重复，请使用下载的模板")
                continue
            for row_no, cells in enumerate(rows, 2):
                if all(c.value in (None, "") for c in cells):
                    continue
                total += 1
                before = len(errors)
                rec = dict(zip(headers, (c.value for c in cells)))
                for column, cell in zip(headers, cells):
                    if cell.data_type == "f":
                        error(sheet, row_no, column, "不支持公式，请先粘贴为值")
                key = "username" if sheet == "用户" else "student_no"
                username, name = _text(rec.get(key)), _text(rec.get("real_name"))
                if not username or not name:
                    error(sheet, row_no, key, "账号/学号和姓名必填")
                if username in seen[sheet]:
                    error(sheet, row_no, key, "文件内账号/学号重复")
                seen[sheet].add(username)
                if len(username) > 100:
                    error(sheet, row_no, key, "账号/学号不能超过 100 个字符")
                org_key = "org_code" if sheet == "用户" else "branch_code"
                org = db.query(Org).filter(Org.code == _text(rec.get(org_key))).first()
                if not org or not org.is_active:
                    error(sheet, row_no, org_key, "组织不存在或已停用")
                elif not can_manage_user(actor, org):
                    error(sheet, row_no, org_key, "无权在该组织下导入")
                elif sheet == "学生档案" and org.org_type != "BRANCH":
                    error(sheet, row_no, org_key, "须填写党支部代码")
                gender = _text(rec.get("gender")) or None
                if gender not in (None, "男", "女"):
                    error(sheet, row_no, "gender", "性别只能为男或女")
                role = (_text(rec.get("role")) or "STUDENT") if sheet == "用户" else "STUDENT"
                if role not in ROLES or (role in PRIVILEGED and actor.role not in PRIVILEGED):
                    error(sheet, row_no, "role", "角色无效或无权授予该角色")
                user = db.query(User).filter(User.username == username).first() if username else None
                planned = user_plans.get(username)
                if user and (not user.is_active or user.real_name != name or not org or user.org_id != org.id):
                    error(sheet, row_no, key, "已有账号的姓名、组织不匹配或账号已停用")
                if planned and (planned["name"] != name or not org or planned["org"].id != org.id or planned["role"] != role):
                    error(sheet, row_no, key, "与用户工作表的姓名、组织或角色不一致")
                batch, birth, nodes, college, existing = selected_batch, None, [], None, None
                if sheet == "学生档案":
                    if org:
                        college = db.get(Org, org.parent_id) if org.parent_id else None
                        if not college or college.org_type != "COLLEGE" or not college.is_active:
                            error(sheet, row_no, org_key, "党支部须隶属于有效学院")
                    if user and not db.query(UserRole.id).filter_by(user_id=user.id, role="STUDENT", org_id=org.id if org else None).first():
                        error(sheet, row_no, key, "已有账号不是该支部的学生，不能自动关联档案")
                    existing = db.query(Archive).filter(Archive.student_no == username).first()
                    linked = db.query(Archive).filter(Archive.student_user_id == user.id).first() if user else None
                    if (existing and (existing.real_name != name or not org or existing.branch_id != org.id
                                      or not user or existing.student_user_id != user.id)) or (linked and linked.student_no != username):
                        error(sheet, row_no, key, "已有档案与本行身份或支部不匹配")
                    batch_name = _text(rec.get("batch_name"))
                    if batch_name:
                        matches = db.query(Batch).filter(Batch.name == batch_name).all()
                        if len(matches) != 1:
                            error(sheet, row_no, "batch_name", "批次名称不存在或不唯一")
                        elif selected_batch and selected_batch.id != matches[0].id:
                            error(sheet, row_no, "batch_name", "与所选建档批次不一致")
                        else:
                            batch = matches[0]
                    if batch and batch.status != 'ACTIVE':
                        error(sheet, row_no, 'batch_name', '批次已关闭，不能导入新档案')
                    try:
                        birth = _date(rec.get("birth_date"))
                        if birth and birth > now().date():
                            raise ValueError("出生日期不能晚于今天")
                    except ValueError as exc:
                        error(sheet, row_no, "birth_date", str(exc))
                    for code, value in rec.items():
                        if not code.startswith("T") or value in (None, ""):
                            continue
                        definition = node_defs.get(code)
                        if not definition or not definition.is_active or definition.value_kind == "DERIVED":
                            error(sheet, row_no, code, "时间节点不存在、已停用或不能直接导入")
                            continue
                        try:
                            node_date = _date(value)
                            if node_date and node_date > now().date():
                                raise ValueError('节点日期不能晚于今天')
                            nodes.append((code, node_date))
                        except ValueError:
                            error(sheet, row_no, code, "节点日期无效，请使用 Excel 日期或 YYYY-MM-DD")
                if len(errors) != before:
                    continue
                valid += 1
                if not user and not planned:
                    user_plans[username] = dict(username=username, name=name, org=org, role=role,
                                                gender=gender, rec=rec, user=None)
                if sheet == "学生档案" and not existing:
                    archive_plans.append(dict(username=username, name=name, org=org, college=college,
                                              batch=batch, birth=birth, nodes=nodes, rec=rec, gender=gender, user=user))
    finally:
        wb.close()
    result = dict(total=total, valid=valid, created_users=0, created_archives=0, imported_nodes=0, errors=errors)
    if dry_run:
        return result
    if errors:
        raise unprocessable("导入存在错误，未写入任何数据", errors)
    try:
        with db.begin_nested():
            passwords = []
            for plan in user_plans.values():
                pwd = random_password()
                rec = plan["rec"]
                user = User(username=plan["username"], password_hash=hash_password(pwd), real_name=plan["name"],
                            org_id=plan["org"].id, gender=plan["gender"], email=_text(rec.get("email")) or None,
                            phone_enc=encrypt_field(_text(rec.get("phone")) or None),
                            phone_masked=mask_phone(_text(rec.get("phone")) or None),
                            must_change_password=True, created_at=now(), updated_at=now())
                db.add(user)
                db.flush()
                db.add(UserRole(user_id=user.id, role=plan["role"], org_id=plan["org"].id, created_at=now()))
                plan["user"] = user
                passwords.append((user.username, pwd))
                result["created_users"] += 1
            for plan in archive_plans:
                user = plan["user"] or user_plans[plan["username"]]["user"]
                rec = plan["rec"]
                archive = Archive(student_user_id=user.id, student_no=plan["username"], real_name=plan["name"],
                    college_id=plan["college"].id, branch_id=plan["org"].id, batch_id=plan["batch"].id if plan["batch"] else None,
                    birth_date=plan["birth"], gender=plan["gender"], ethnicity=rec.get("ethnicity"),
                    native_place=rec.get("native_place"), major=rec.get("major"), class_name=rec.get("class_name"),
                    current_stage=1, status="ACTIVE", anon_code="A" + uuid4().hex, data_source="IMPORT",
                    created_by=actor.id, created_at=now(), updated_at=now())
                db.add(archive)
                db.flush()
                db.add(Submission(archive_id=archive.id, stage_code=1, batch_id=archive.batch_id,
                    college_id=archive.college_id, branch_id=archive.branch_id, status="DRAFT", round_no=1,
                    ai_status="NONE", created_at=now(), updated_at=now()))
                for code, value in plan["nodes"]:
                    db.add(TimeNodeExtraction(archive_id=archive.id, node_code=code, value_date=value,
                                             method="IMPORT", is_current=True, created_at=now()))
                db.flush()
                merge_archive_nodes(db, archive.id)
                result["created_archives"] += 1
                result["imported_nodes"] += len(plan["nodes"])
            if passwords:
                from services.import_passwords import create_password_file
                result["initial_passwords_file"] = create_password_file(db, actor, passwords)
            db.flush()
    except IntegrityError:
        raise unprocessable("导入数据与已有记录冲突，请重新预校验后重试")
    return result
