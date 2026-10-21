"""Statistics and real XLSX regressions in an isolated PostgreSQL schema."""
from datetime import datetime
from io import BytesIO
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook
from sqlalchemy import insert

from errors import AppError
from models import Archive, ExportFile, Issue, Material, MaterialType, ReviewRecord, Submission, User
from services.exporter import run_export_job
from services.stats import ai_quality, normalize_filters, overview, top_issues
from utils import TZ8, now
from tests.test_review_regressions import test_engine, case, as_client  # noqa: F401


def issue(c, material=None, status="CONFIRMED", sub=None, source="AI_RULE"):
    return c.add(Issue, submission_id=(sub or c.sub).id, report_id=c.report.id,
                 material_id=material.id if material else None, source=source, category="FORMAT",
                 severity="WARN", rule_code="R-F01", title="签名", description="待核对",
                 review_status=status, created_at=now(), updated_at=now())


def second_sub(c, **values):
    return c.add(Submission, archive_id=c.arch.id, stage_code=2, batch_id=c.batches[0].id,
                 college_id=c.college.id, branch_id=c.branches[0].id,
                 created_at=now(), updated_at=now(), **values)


def export_book(c, monkeypatch, tmp_path, export_type, filters=None, scope=None):
    from services import storage
    monkeypatch.setattr(storage, "abs_path", lambda rel: tmp_path / rel)
    job = SimpleNamespace(id=123, created_by=c.org.id, params={"export_type": export_type,
        "filters": filters or {}, "scope_college_id": scope}, result=None)
    run_export_job(c.db, job)
    record = c.db.get(ExportFile, job.result["export_id"])
    return job, load_workbook(BytesIO((tmp_path / record.stored_path).read_bytes()), data_only=True)


def test_material_types_aggregate_distinct_materials_and_pending_does_not_dilute_precision(case):
    c = case
    mt = c.db.query(MaterialType).first()
    for status in ("CONFIRMED", "REJECTED", "PENDING"):
        m = c.add(Material, archive_id=c.arch.id, submission_id=c.sub.id, stage_code=1,
                  title="材料", type_id=mt.id, status="ACTIVE", created_at=now(), updated_at=now())
        issue(c, m, status)
    grouped = top_issues(c.db, c.org, group_by="material_type")
    assert [(g["key"], g["count"]) for g in grouped] == [(mt.code, 3)]
    quality = ai_quality(c.db, c.org)
    assert quality["by_rule"][0]["precision"] == .5
    assert quality["by_material_type"][0]["name"] == mt.name
    assert quality["by_material_type"][0]["total"] == 3
    assert quality["first_pass_rate"] is None


def test_first_pass_uses_completed_packages_and_return_history(case):
    c = case
    c.sub.status = "FILED"
    second_sub(c, status="FILED", round_no=1)
    c.add(ReviewRecord, submission_id=c.sub.id, level="BRANCH", action="RETURN", from_status="BRANCH_REVIEW",
          to_status="RETURNED", reviewer_name="审核员", reviewer_role="BRANCH_REVIEWER", round_no=1,
          created_at=now())
    c.db.flush()
    quality = ai_quality(c.db, c.org)
    assert quality["first_pass_rate"] == .5
    assert quality["completed_submissions"] == 2
    assert quality["first_pass_count"] == 1


def test_ai_quality_dates_are_inclusive_in_china_timezone_and_reach_api(case, as_client):
    c = case
    c.sub.submitted_at = datetime(2026, 3, 1, 23, 59, tzinfo=TZ8)
    later = second_sub(c, submitted_at=datetime(2026, 3, 2, 0, 0, tzinfo=TZ8))
    issue(c)
    issue(c, sub=later, status="REJECTED")
    c.db.flush()
    quality = ai_quality(c.db, c.org, date_from="2026-03-01", date_to="2026-03-01")
    assert quality["ai_issues"] == 1
    assert quality["reports"] == 1
    response = as_client(c.org).get("/api/stats/ai-quality?date_from=2026-03-02")
    assert response.status_code == 200, response.text
    assert response.json()["data"]["ai_issues"] == 1
    assert response.json()["data"]["rejected"] == 1


@pytest.mark.parametrize("filters", [{"date_from": "2026-99-01"}, {"date_from": "2026-03-02", "date_to": "2026-03-01"},
    {"batch_id": -1}, {"stage_code": 6}, {"branch_id": True}])
def test_invalid_filters_rejected(filters):
    with pytest.raises(AppError):
        normalize_filters(filters)


@pytest.mark.parametrize("export_type", ["STATS_REPORT", "DESENSITIZED_DATA", "AI_QUALITY"])
def test_all_exports_apply_stage_date_and_scope(case, monkeypatch, tmp_path, export_type):
    c = case
    c.sub.submitted_at = datetime(2026, 3, 1, 12, tzinfo=TZ8)
    other = second_sub(c, submitted_at=datetime(2026, 4, 1, 12, tzinfo=TZ8))
    issue(c)
    issue(c, sub=other, status="REJECTED")
    c.db.flush()
    filters = {"stage_code": 1, "date_from": "2026-03-01", "date_to": "2026-03-01",
               "college_id": c.college.id + 999}
    job, book = export_book(c, monkeypatch, tmp_path, export_type, filters, c.college.id)
    assert job.result["row_count"] == 1
    assert "筛选说明" in book.sheetnames
    rows = list(book.worksheets[0].values)
    assert len(rows) == 2
    if export_type == "AI_QUALITY":
        assert rows[1][4:7] == (1, 0, 0)
        assert "AI效能汇总" in book.sheetnames
    if export_type == "DESENSITIZED_DATA":
        assert rows[1][0] == c.arch.anon_code
        assert c.arch.real_name not in repr(list(book.worksheets[0].values))
    _, empty = export_book(c, monkeypatch, tmp_path, export_type, filters, c.college.id + 999)
    assert len(list(empty.worksheets[0].values)) == 1


def test_overview_archive_count_obeys_dates(case):
    c = case
    c.sub.submitted_at = datetime(2026, 3, 1, 12, tzinfo=TZ8)
    c.db.flush()
    result = overview(c.db, c.org, date_from="2026-04-01")
    assert result["archives"] == result["submissions"] == 0


def test_export_includes_more_than_5000_packages(case, monkeypatch, tmp_path):
    c = case
    count = 5001
    timestamp = now()
    users = c.db.execute(insert(User).returning(User.id), [dict(username=f"bulk-{i}", password_hash="unused",
        real_name="合成数据", org_id=c.branches[0].id, created_at=timestamp, updated_at=timestamp)
        for i in range(count)]).scalars().all()
    archives = c.db.execute(insert(Archive).returning(Archive.id), [dict(student_user_id=uid, student_no=f"bulk-{i}",
        real_name="合成数据", college_id=c.college.id, branch_id=c.branches[0].id, current_stage=1,
        status="ACTIVE", anon_code=f"anon-{i}", created_at=timestamp, updated_at=timestamp)
        for i, uid in enumerate(users)]).scalars().all()
    c.db.execute(insert(Submission), [dict(archive_id=aid, stage_code=1, college_id=c.college.id,
        branch_id=c.branches[0].id, batch_id=c.batches[1].id, created_at=timestamp, updated_at=timestamp)
        for aid in archives])
    job, book = export_book(c, monkeypatch, tmp_path, "STATS_REPORT", {"batch_id": c.batches[1].id})
    assert job.result["row_count"] == count
    assert len(list(book.worksheets[0].values)) == count + 1


def test_export_api_rejects_bad_filter_and_unknown_type(case, as_client):
    client = as_client(case.org)
    for body in ({"export_type": "typo"}, {"filters": {"date_from": "wrong"}}, {"filters": [1]}):
        assert client.post("/api/stats/export", json=body).status_code == 422
